import openpyxl
import json
import re

def parse_bast(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    
    # 1. Parse Rekap
    rekap_ws = wb['Rekap']
    rekap_info = {}
    title_row = rekap_ws.cell(1, 5).value or rekap_ws.cell(1, 1).value or "Metadata BAST DTSEN"
    rekap_info['title'] = str(title_row).strip()
    
    # Version detection
    version_match = re.search(r'Versi\s*([\d\.\/\w]+)', str(title_row), re.IGNORECASE)
    rekap_info['version'] = version_match.group(1) if version_match else "3/2026"
    rekap_info['source'] = "BPS-Bappenas"
    
    datasets = {
        'keluarga': {'title': 'Set Data Keluarga', 'sheet': 'Set Data Keluarga', 'variables': []},
        'anggota_keluarga': {'title': 'Set Data Anggota Keluarga', 'sheet': 'Set Data Anggota Keluarga', 'variables': []}
    }
    
    for ds_key, ds_meta in datasets.items():
        ws = wb[ds_meta['sheet']]
        rows = list(ws.iter_rows(values_only=True))
        
        # Header is usually row 2 (index 1), data starts from row 4 (index 3)
        for r_idx in range(3, len(rows)):
            row = rows[r_idx]
            if not row or row[0] is None:
                continue
            no_val = str(row[0]).strip()
            if no_val.startswith('-') or not no_val.replace('.', '').isdigit():
                continue
            
            no_num = int(float(no_val))
            label = str(row[1] or '').strip().replace('\n', ' ')
            col_name = str(row[2] or '').strip().lower()
            definition = str(row[3] or '').strip()
            datatype = str(row[4] or '').strip().lower()
            keterangan = str(row[5] or '').strip()
            
            if not col_name:
                continue
                
            # Parse length from datatype e.g. varchar(16), varchar(100), datetime(yyyy-mm-dd)
            length_match = re.search(r'\((\d+)\)', datatype)
            max_len = int(length_match.group(1)) if length_match else None
            
            # Parse allowed categorical values/codes from keterangan
            # e.g. "1. Laki-laki\n2. Perempuan" or "1. Penerima\n0. Tidak"
            allowed_codes = []
            if keterangan:
                # Look for lines like "1. Option" or "0. Option"
                code_matches = re.findall(r'(?:^|\n)\s*(\d+)\.\s*([^\n\r]+)', keterangan)
                for code_num, code_label in code_matches:
                    allowed_codes.append({
                        'code': code_num.strip(),
                        'label': code_label.strip()
                    })
            
            # Determine validation rules
            rules = {
                'required': col_name in ('nomor_induk_kependudukan', 'nomor_kartu_keluarga', 'nama'),
                'max_length': max_len,
                'datatype': datatype,
                'allowed_codes': allowed_codes,
            }
            
            if col_name in ('nomor_induk_kependudukan', 'nomor_kartu_keluarga'):
                rules['exact_length'] = 16
                rules['regex'] = '^[0-9]{16}$'
                rules['error_type'] = 'Critical'
            elif 'desil' in col_name:
                rules['range'] = [1, 10]
                rules['error_type'] = 'Critical'
            elif 'usia' in col_name or 'umur' in col_name:
                rules['range'] = [0, 120]
                rules['error_type'] = 'Critical'
            elif 'nama' in col_name:
                rules['regex'] = "^[a-zA-Z\\s\\.\\,\\'\\-]+$"
                rules['error_type'] = 'Critical'
            elif 'tanggal' in col_name or 'datetime' in datatype:
                rules['format'] = 'date'
                rules['error_type'] = 'Warning'
            elif allowed_codes:
                rules['error_type'] = 'Warning'
            else:
                rules['error_type'] = 'Warning'
                
            var_item = {
                'no': no_num,
                'key': col_name,
                'label': label,
                'definition': definition,
                'datatype': datatype,
                'keterangan': keterangan,
                'rules': rules
            }
            ds_meta['variables'].append(var_item)
            
    result = {
        'metadata_info': rekap_info,
        'datasets': datasets,
        'total_variables': len(datasets['keluarga']['variables']) + len(datasets['anggota_keluarga']['variables']),
        'total_keluarga_vars': len(datasets['keluarga']['variables']),
        'total_anggota_vars': len(datasets['anggota_keluarga']['variables']),
        'parsed_at': '2026-10-07 09:10:00'
    }
    return result

if __name__ == '__main__':
    res = parse_bast('storage/app/bast_metadata/Versi_3_2026_Lampiran_BAST_Metadata_DTSEN.xlsx')
    print("Version:", res['metadata_info']['version'])
    print("Title:", res['metadata_info']['title'])
    print(f"Keluarga Vars: {res['total_keluarga_vars']}, Anggota Vars: {res['total_anggota_vars']}, Total: {res['total_variables']}")
    
    with open('storage/app/bast_metadata/active_rules.json', 'w', encoding='utf-8') as f:
        json.dump(res, f, indent=2, ensure_ascii=False)
    print("Saved to storage/app/bast_metadata/active_rules.json successfully!")
