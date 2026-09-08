import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = 'Arial'
    if level == 1:
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
    elif level == 2:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(53, 89, 143) # Slate Blue
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(40, 40, 40)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    return h

def build_word_document():
    doc = Document()
    
    # Set Margins (1 inch / 2.54 cm all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Global Style Normal Paragraph
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Arial'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 30, 30)

    # ---------------------------------------------------------
    # COVER / TITLE BLOCK
    # ---------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("DOKUMENTASI LENGKAP SISTEM PADU v1.02")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(16, 44, 87)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Pengolah dan Analisis Data Terpadu — DTSEN 2026 High-Speed Analytics Engine")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    # Divider line table
    table_div = doc.add_table(rows=1, cols=1)
    table_div.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_div = table_div.cell(0, 0)
    cell_div.width = Inches(6.5)
    set_cell_background(cell_div, "102C57")
    p_div = cell_div.paragraphs[0]
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(2)

    # Doc Metadata
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(18)
    p_meta.paragraph_format.space_after = Pt(24)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Versi Dokumen: 1.02 (Updated) | Tanggal Rilis: 2026 | Lingkungan: 100% Offline Windows")
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_page_break()

    # ---------------------------------------------------------
    # BAB 1: GAMBARAN UMUM PROYEK & SPESIFIKASI SISTEM
    # ---------------------------------------------------------
    add_styled_heading(doc, "BAB 1: GAMBARAN UMUM PROYEK & SPESIFIKASI SISTEM", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Aplikasi ").bold = True
    p.add_run("PADU (Pengolah dan Analisis Data Terpadu) v1.02 ").bold = True
    p.add_run("adalah perangkat lunak analitik data kependudukan dan sosial-ekonomi berbasis web yang dirancang khusus untuk mengelola dataset skala besar ")
    p.add_run("DTSEN (Data Terpadu Sosial Ekonomi Nasional / Registrasi Sosial Ekonomi 2026) ").bold = True
    p.add_run("dengan kapasitas hingga ")
    p.add_run("13+ juta baris data (setara > 2.7 GB dataset CSV)").bold = True
    p.add_run(" secara seratus persen offline di perangkat pengguna.")

    add_styled_heading(doc, "1.1 Tujuan dan Latar Belakang", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Tujuan utama dari pengoperasian PADU v1.02 adalah menyediakannya alat bantu analisis data cepat untuk auditor, pengambil kebijakan, dan operator di lapangan agar dapat:\n")
    p.add_run("1. Melakukan pengolahan data raksasa (13M+ baris) secara instan tanpa memerlukan server database terpusat yang rumit.\n")
    p.add_run("2. Memeriksa tingkat validitas dan keutuhan data (Quality Score Audit) secara otomatis.\n")
    p.add_run("3. Melindungi data identitas pribadi penduduk (PII Data Masking) saat data ditayangkan di layar antarmuka.\n")
    p.add_run("4. Memfasilitasi pencarian NIK/KK, pemfilteran dinamis multi-variabel, serta agregasi statistik metrik (KPI) secara real-time.")

    add_styled_heading(doc, "1.2 Persyaratan Hardware & Lingkungan Eksekusi", level=2)
    
    # Table Hardware
    t_hw = doc.add_table(rows=5, cols=3)
    t_hw.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_hw = ["Komponen", "Spesifikasi Minimal", "Keterangan & Catatan Opsional"]
    for i, h_text in enumerate(headers_hw):
        cell = t_hw.cell(0, i)
        set_cell_background(cell, "102C57")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 120, 120, 150, 150)

    rows_data_hw = [
        ["Sistem Operasi", "Windows 10 / Windows 11 (64-bit)", "Mendukung arsitektur Windows x64 modern."],
        ["Memori (RAM)", "Minimal 4 GB RAM", "Alokasi memori sangat efisien berkat DuckDB Vectorized Pages."],
        ["Penyimpanan (Disk)", "Minimal 2 GB Ruang Kosong", "Tergantung ukuran dataset CSV yang diimpor ke sistem."],
        ["Runtime Bundled", "Portable PHP 8.2 & Python 3.11", "Sudah terikut otomatis di dalam folder tanpa instalasi admin."]
    ]

    for row_idx, r_data in enumerate(rows_data_hw, start=1):
        bg_color = "F8F9FA" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(r_data):
            cell = t_hw.cell(row_idx, col_idx)
            set_cell_background(cell, bg_color)
            p = cell.paragraphs[0]
            p.add_run(text)
            set_cell_margins(cell, 100, 100, 150, 150)

    # ---------------------------------------------------------
    # BAB 2: ARSITEKTUR KEAMANAN & PERLINDUNGAN PRIVASI DATA
    # ---------------------------------------------------------
    add_styled_heading(doc, "BAB 2: ARSITEKTUR KEAMANAN & PERLINDUNGAN PRIVASI DATA", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Aspek keamanan pada aplikasi PADU v1.02 berfokus pada dua pilar utama: ")
    p.add_run("Perlindungan Kerahasiaan Data Pribadi (PII Data Masking) ").bold = True
    p.add_run("dan ")
    p.add_run("Isolasi Lingkungan 100% Offline (Zero-Cloud Leakage).").bold = True

    add_styled_heading(doc, "2.1 Penyamaran Data Sensitif (PII Data Masking)", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Untuk mencegah kebocoran informasi identitas kependudukan saat layar dipresentasikan atau diakses operator, seluruh variabel sensitif disamarkan secara otomatis oleh ")
    p.add_run("App\\Services\\DataMaskingService").bold = True
    p.add_run(" dengan aturan sebagai berikut:")

    # Table Masking Rules
    t_mask = doc.add_table(rows=8, cols=4)
    t_mask.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_mask = ["Variabel Data", "Contoh Data Asli", "Hasil Masking Layar", "Aturan & Algoritma Penyamaran"]
    for i, h_text in enumerate(headers_mask):
        cell = t_mask.cell(0, i)
        set_cell_background(cell, "35598F")
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 120, 120, 150, 150)

    mask_data = [
        ["NIK (Nomor Induk Kependudukan)", "3201021508900001", "3201************", "Menampilkan 4 digit wilayah awal, 12 digit disensor *"],
        ["Nama Lengkap", "Budi Santoso", "B*** S******", "Huruf pertama tetap, sisa karakter tiap kata disensor *"],
        ["Tanggal Lahir", "1992-05-14", "1992-**-**", "Hanya menampilkan 4 digit tahun lahir"],
        ["Alamat Rumah", "Jl. Sudirman No. 45", "Jl. S******* No. **", "Menyamarkan nama jalan dan nomor bangunan"],
        ["Nominal Gaji / Pendapatan", "Rp 8.500.000", "Rp 8.xxx.xxx", "Menyamarkan digit ribuan dengan karakter x"],
        ["RT / RW", "RT 003 / RW 005", "RT *** / RW ***", "Mengganti seluruh digit angka RT/RW menjadi *"],
        ["Email", "budi.santoso@gmail.com", "b***@gmail.com", "Menampilkan karakter pertama username dan domain"]
    ]

    for row_idx, r_data in enumerate(mask_data, start=1):
        bg_color = "F8F9FA" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(r_data):
            cell = t_mask.cell(row_idx, col_idx)
            set_cell_background(cell, bg_color)
            p = cell.paragraphs[0]
            p.add_run(text)
            set_cell_margins(cell, 100, 100, 150, 150)

    add_styled_heading(doc, "2.2 Isolasi 100% Offline & Anti-Leakage Architecture", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("1. Zero Outbound Telemetry: ").bold = True
    p.add_run("Sistem tidak mengirimkan sinyal analitik, log crash, atau HTTP request ke internet.\n")
    p.add_run("2. Local In-Process Storage: ").bold = True
    p.add_run("Seluruh data tersimpan secara lokal pada berkas database `database/dataset.duckdb`.\n")
    p.add_run("3. Application Key Encryption: ").bold = True
    p.add_run("Setiap instalasi menggunakan Laravel APP_KEY unik yang diproduksi otomatis oleh `artisan key:generate`.\n")
    p.add_run("4. Parameterized Query Protection: ").bold = True
    p.add_run("Seluruh eksekusi pencarian dan filter pada engine Python DuckDB menggunakan *parameterized filtering* untuk mencegah serangan SQL Injection.")

    # ---------------------------------------------------------
    # BAB 3: ARSITEKTUR & CARA KERJA SISTEM
    # ---------------------------------------------------------
    add_styled_heading(doc, "BAB 3: ARSITEKTUR & CARA KERJA SISTEM", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("PADU v1.02 menggabungkan keunggulan antarmuka web modern dengan keahlian engine pemrosesan data OLAP dalam arsitektur ")
    p.add_run("Hybrid Dual-Engine (Laravel 12 + Python DuckDB)").bold = True
    p.add_run(".")

    add_styled_heading(doc, "3.1 Alur Kerja Pipeline Data (Data Pipeline Flow)", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("1. Peluncuran Sistem: ").bold = True
    p.add_run("Pengguna menjalankan `JALANKAN_PADU.bat`. Script akan mengaktifkan server Portable PHP pada Port 8000 dan membuka browser otomatis.\n")
    p.add_run("2. Ingesti Data (Upload & Import): ").bold = True
    p.add_run("File CSV/XLSX diunggah melalui antarmuka web. Controller `DtsenController::importLocalPath` memanggil script `scratch/fast_import.py` secara background subprocess.\n")
    p.add_run("3. 1-Pass Vectorized Ingestion: ").bold = True
    p.add_run("Script Python memanfaatkan `duckdb.read_csv_auto` untuk menguraikan 13+ juta baris data langsung ke tabel `individus` di `database/dataset.duckdb` dalam waktu < 3 detik.\n")
    p.add_run("4. Evaluasi Kualitas Real-time: ").bold = True
    p.add_run("Dalam 1-pass query yang sama, fungsi SQL `ARRAY_FILTER` memeriksa NIK, KK, Nama, Usia, Gaji, dan Desil, lalu menentukan status `Valid` atau `Critical`.\n")
    p.add_run("5. Agregasi statistik HyperLogLog: ").bold = True
    p.add_run("Sistem menghitung total baris dan total KK unik secara instan (0.01 detik) menggunakan algoritma HyperLogLog probabilistic counting.\n")
    p.add_run("6. Penyajian Data & Ekspor: ").bold = True
    p.add_run("Data disajikan ke UI Blade melalui pagination `fast_duckdb.py`. Pengguna dapat mengunduh paket data ZIP yang berisi CSV bersih dan file audit `metadata.txt`.")

    add_styled_heading(doc, "3.2 Aturan Validasi Quality Check Engine", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Kriteria penentuan status kualitas data ditentukan sebagai berikut:\n")
    p.add_run("• CRITICAL: ").bold = True
    p.add_run("Nama kosong/mengandung angka/simbol khusus, NIK/KK != 16 digit, Desil di luar 1-10, Usia < 0 atau > 120 tahun, Gaji < 0.\n")
    p.add_run("• WARNING: ").bold = True
    p.add_run("Nilai variabel atribut pendukung (seperti jenis lantai, dinding, status kerja) kosong/missing value.\n")
    p.add_run("• VALID: ").bold = True
    p.add_run("Seluruh syarat integritas dan kelengkapan variabel terpenuhi.")

    # ---------------------------------------------------------
    # BAB 4: PANDUAN MODIFIKASI & PENGGANTIAN SUB-SISTEM
    # ---------------------------------------------------------
    add_styled_heading(doc, "BAB 4: PANDUAN MODIFIKASI & PENGGANTIAN SUB-SISTEM", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Dokumentasi ini memberikan panduan rinci bagi pengembang yang ingin mengubah atau menyesuaikan komponen sub-sistem PADU v1.02.")

    # Table Modification Guide
    t_mod = doc.add_table(rows=7, cols=4)
    t_mod.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_mod = ["Komponen yang Ingin Diubah", "Lokasi Berkas / File Target", "Fungsi / Method Terkait", "Petunjuk & Cara Modifikasi"]
    for i, h_text in enumerate(headers_mod):
        cell = t_mod.cell(0, i)
        set_cell_background(cell, "102C57")
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 120, 120, 150, 150)

    mod_data = [
        ["Aturan Masking Data (PII)", "app/Services/DataMaskingService.php", "maskNik(), maskNama(), maskGaji(), maskAlamat()", "Edit ekspresi reguler (regex) atau jumlah karakter sensor pada method target."],
        ["Aturan Quality Check (PHP)", "app/Services/DataQualityCheckService.php", "evaluate()", "Tambahkan kondisi validasi variabel baru pada blok evaluasi Critical/Warning."],
        ["Aturan Quality Check (DuckDB)", "scratch/fast_import.py", "run_fast_import()", "Ubah ekspresi SQL ARRAY_FILTER pada pembuatan tabel individus di DuckDB."],
        ["Engine Database / Analytics", "scratch/fast_duckdb.py", "run_duckdb_query()", "Modifikasi fungsi build_where_clause() atau logika kalkulasi mode kpi_metrics."],
        ["Antarmuka UI Web", "resources/views/dtsen/index.blade.php", "Blade Template & Alpine.js", "Sesuaikan struktur HTML, komponen filter multi-checkbox, atau tabel pratinjau."],
        ["Peluncur Batch Windows", "JALANKAN_PADU.bat", "Script Command Batch", "Ubah alokasi port server (misal --port=8080) atau mekanisme deteksi path PHP."]
    ]

    for row_idx, r_data in enumerate(mod_data, start=1):
        bg_color = "F8F9FA" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(r_data):
            cell = t_mod.cell(row_idx, col_idx)
            set_cell_background(cell, bg_color)
            p = cell.paragraphs[0]
            p.add_run(text)
            set_cell_margins(cell, 100, 100, 150, 150)

    # ---------------------------------------------------------
    # BAB 5: PEMETAAN FUNGSI, METHOD, & SERVICE
    # ---------------------------------------------------------
    add_styled_heading(doc, "BAB 5: PEMETAAN FUNGSI, METHOD, & SERVICE", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Berikut adalah rincian fungsi dan method utama yang terlibat dalam sistem PADU v1.02:")

    # Table Functions
    t_fn = doc.add_table(rows=10, cols=3)
    t_fn.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_fn = ["Nama File / Class", "Nama Fungsi / Method", "Deskripsi & Peran Teknis Fungsi"]
    for i, h_text in enumerate(headers_fn):
        cell = t_fn.cell(0, i)
        set_cell_background(cell, "35598F")
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 120, 120, 150, 150)

    fn_data = [
        ["DtsenController.php", "index()", "Menampilkan halaman utama analitik, memicu query data halaman & agregasi statistik KPI."],
        ["DtsenController.php", "importLocalPath()", "Menerima path file CSV/XLSX lokal dan memanggil subprocess fast_import.py."],
        ["DtsenController.php", "exportCsv()", "Memicu subprocess fast_export.py untuk membentuk berkas ZIP paket data."],
        ["DataMaskingService.php", "maskNik(), maskNama()", "Fungsi statis penyensoran NIK 16 digit dan Nama Lengkap."],
        ["DataQualityCheckService.php", "evaluate()", "Mengevaluasi keabsahan data individu & keluarga untuk menghasilkan status kualitas."],
        ["fast_import.py", "run_fast_import()", "Ingesti cepat 1-pass CSV ke DuckDB, membentuk tabel individus & menghitung stats."],
        ["fast_duckdb.py", "build_where_clause()", "Membangun klausa SQL WHERE dinamis dari pencarian & multi-filter user."],
        ["fast_duckdb.py", "run_duckdb_query()", "Menjalankan mode query: page_data, total_system_rows, distinct_values, kpi_metrics."],
        ["fast_export.py", "run_fast_export()", "Ekspor data cepat dari DuckDB ke CSV, pembentukan metadata.txt, & zipping paket."]
    ]

    for row_idx, r_data in enumerate(fn_data, start=1):
        bg_color = "F8F9FA" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(r_data):
            cell = t_fn.cell(row_idx, col_idx)
            set_cell_background(cell, bg_color)
            p = cell.paragraphs[0]
            p.add_run(text)
            set_cell_margins(cell, 100, 100, 150, 150)

    # ---------------------------------------------------------
    # BAB 6: PEMETAAN BERKAS LENGKAP PROYEK
    # ---------------------------------------------------------
    add_styled_heading(doc, "BAB 6: PEMETAAN BERKAS LENGKAP PROYEK", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Daftar seluruh berkas utama yang terlibat dalam pengoperasian aplikasi PADU v1.02:")

    # Table Files
    t_file = doc.add_table(rows=11, cols=3)
    t_file.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_file = ["Jalur Berkas (File Path)", "Peran & Deskripsi Berkas", "Kategori Komponen"]
    for i, h_text in enumerate(headers_file):
        cell = t_file.cell(0, i)
        set_cell_background(cell, "102C57")
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 120, 120, 150, 150)

    file_data = [
        ["JALANKAN_PADU.bat", "Peluncur batch zero-config Windows untuk memulai server PHP 8000 & browser.", "Launcher Script"],
        ["SRS_PADU_v1.02.md", "Dokumen Software Requirements Specification berstandar ISO/IEEE.", "Documentation"],
        ["README.md", "Panduan proyek, arsitektur, instalasi, dan petunjuk modifikasi pengembang.", "Documentation"],
        ["app/Http/Controllers/DtsenController.php", "Controller utama pengatur rute HTTP, filter, import, dan export.", "Backend Controller"],
        ["app/Services/DataMaskingService.php", "Service penyamaran PII data sensitif (NIK, Nama, Gaji, Alamat).", "Security Service"],
        ["app/Services/DataQualityCheckService.php", "Service aturan audit kualitas data (Valid, Warning, Critical).", "Quality Audit Service"],
        ["scratch/fast_import.py", "Engine Python ingesti cepat 1-pass CSV ke DuckDB.", "Python Analytics Engine"],
        ["scratch/fast_duckdb.py", "Engine Python query terurai, multi-filter, & agregasi KPI DuckDB.", "Python Analytics Engine"],
        ["scratch/fast_export.py", "Engine Python ekspor data cepat ke CSV dan pembuatan ZIP audit log.", "Python Export Engine"],
        ["resources/views/dtsen/index.blade.php", "Antarmuka antarmuka utama pengguna (Dashboard UI Blade & Alpine.js).", "Frontend View UI"]
    ]

    for row_idx, r_data in enumerate(file_data, start=1):
        bg_color = "F8F9FA" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(r_data):
            cell = t_file.cell(row_idx, col_idx)
            set_cell_background(cell, bg_color)
            p = cell.paragraphs[0]
            p.add_run(text)
            set_cell_margins(cell, 100, 100, 150, 150)

    # Save Document
    out_path = os.path.join(os.getcwd(), "Dokumentasi_Lengkap_Sistem_PADU_v1.02_Updated.docx")
    doc.save(out_path)
    print(f"[SUCCESS] Dokumen Word berhasil dibuat: {out_path}")

if __name__ == '__main__':
    build_word_document()
