import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    """Sets background color of a cell (hex string e.g. '1F4E78')"""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa (1 pt = 20 dxa)"""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    if level == 1:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78) # Dark Blue
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6) # Steel Blue
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
    return p

def add_paragraph_styled(doc, text="", bold_prefix=None, italic=False, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    if text:
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(11)
        r_text.font.italic = italic
        r_text.font.bold = bold
        r_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def build_srs_doc(output_path):
    doc = docx.Document()
    
    # Page Setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Title Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run("DRAF SPESIFIKASI REQUIREMENT SRS\nKETERANGAN APLIKASI & SPESIFIKASI LENGKAP USE CASE (UC-01 s/d UC-10)")
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    run_s = p_sub.add_run("Aplikasi PADU v1.02 — Pengolah & Analisis Data Terpadu (DTSEN 2026 Analytics Engine)\nDokumen Acuan & Referensi Penyusunan SRS untuk Tim / Rekan Kerja")
    run_s.font.name = 'Calibri'
    run_s.font.size = Pt(11)
    run_s.font.italic = True
    run_s.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # BAB 1: KETERANGAN APLIKASI & ARSITEKTUR
    # ==========================================
    add_heading_styled(doc, "1. KETERANGAN APLIKASI & ARSITEKTUR SISTEM", level=1)
    
    add_heading_styled(doc, "1.1 Deskripsi & Pengantar Sistem", level=2)
    add_paragraph_styled(doc, "PADU v1.02 (Pengolah dan Analisis Data Terpadu — DTSEN 2026 Analytics Engine)", bold_prefix="• Nama Sistem: ")
    add_paragraph_styled(doc, "Aplikasi web analitik performa tinggi untuk mengimpor, mengaudit kualitas data (quality check), melakukan pencarian/filtering instan, serta menyamarkan data sensitif (data masking) pada dataset DTSEN (Data Terpadu Sosial Ekonomi Nasional 2026) skala besar (hingga 13+ juta baris data) secara 100% Offline (tanpa koneksi internet/cloud).", bold_prefix="• Tujuan Utama: ")

    add_heading_styled(doc, "1.2 Arsitektur Dual-Engine", level=2)
    add_paragraph_styled(doc, "Laravel 12 (PHP 8.2 Portable) yang responsif untuk antarmuka web, routing, controller, dan manajemen navigasi pengguna.", bold_prefix="• Frontend / Web Interface: ")
    add_paragraph_styled(doc, "Python 3.11 + DuckDB (Vectorized OLAP Execution Engine) yang bertugas memproses jutaan baris data secara ultra-cepat (kecepatan ingesti > 1,5 juta baris/detik dan respons filter < 0,05 detik).", bold_prefix="• High-Speed Processing Engine: ")

    add_heading_styled(doc, "1.3 Lingkungan Operasi & Persyaratan Sistem", level=2)
    add_paragraph_styled(doc, "Windows 10 / Windows 11 (64-bit).", bold_prefix="• Sistem Operasi: ")
    add_paragraph_styled(doc, "Minimal 4 GB RAM (Penggunaan alokasi memori sangat efisien via DuckDB memory-mapped vector pages).", bold_prefix="• Memori (RAM): ")
    add_paragraph_styled(doc, "Minimal 2 GB ruang kosong.", bold_prefix="• Penyimpanan Disk: ")
    add_paragraph_styled(doc, "Standalone Web Server pada Port 8000 (http://127.0.0.1:8000/).", bold_prefix="• Port Server: ")

    add_heading_styled(doc, "1.4 Fitur-Fitur Utama Sistem (System Core Features)", level=2)
    add_paragraph_styled(doc, "Membaca file CSV/XLSX skala besar (hingga 13M+ baris). Melakukan pemetaan header otomatis (synonym mapping) untuk 48 variabel individu dan 52 variabel keluarga standar BPS/Bappenas. Pemrosesan impor bertahap (chunked background ingestion) lengkap dengan tombol Cancel dan Progress Bar.", bold_prefix="1. Vectorized Data Ingestion (Import Engine): ")
    add_paragraph_styled(doc, "Mengevaluasi keabsahan data secara otomatis saat diimpor dan membaginya ke 3 status:\n  - 🔴 Critical: Kesalahan fatal (NIK/KK tidak 16 digit, Nama memuat angka/karakter aneh, Desil < 1 atau > 10, Umur < 0 atau > 120 tahun, Gaji negatif).\n  - 🟡 Warning: Terdapat data opsional yang kosong (missing values).\n  - 🟢 Valid: Seluruh aturan validasi terpenuhi.", bold_prefix="2. Quality Check & Data Audit Engine: ")
    add_paragraph_styled(doc, "Menyamarkan data pribadi sensitif (Personally Identifiable Information) pada UI preview secara otomatis:\n  - NIK: 3201021508900001 ➔ 3201************\n  - Nama: Budi Santoso ➔ B*** S******\n  - Tanggal Lahir: 1992-05-14 ➔ 1992-**-**\n  - Gaji: Rp 8.500.000 ➔ Rp 8.xxx.xxx\n  - Alamat/RT/RW: Disamarkan secara otomatis.", bold_prefix="3. Data Masking & PII Security Protection: ")
    add_paragraph_styled(doc, "Menghitung statistik (Total Baris, Total KK, Rata-rata/Max/Min Gaji, Distribusi Desil 1–10, Jumlah PBI) dalam hitungan milidetik (< 0,05 detik).", bold_prefix="4. Instant Aggregation & KPI Analytics: ")
    add_paragraph_styled(doc, "Pencarian cepat berdasarkan NIK, KK, Nama, serta kombinasi multi-filter (Desil, Rentang Usia, Rentang Gaji, Status Kerja, Wilayah/Provinsi/Kabupaten).", bold_prefix="5. Dynamic Multi-Column Filtering & Search: ")
    add_paragraph_styled(doc, "Export Data Bersih: Menghasilkan file terkompresi .zip berisi dataset data_dtsen.csv dan log audit metadata.txt.\nExport Data Error: Menghasilkan file .csv yang hanya berisi baris bertanda Critical untuk ditindaklanjuti/diperbaiki oleh petugas lapangan.", bold_prefix="6. Export Engine (ZIP Package & Error CSV): ")
    add_paragraph_styled(doc, "Fitur pantau log eksekusi Python/DuckDB, unduh log audit, dan reset/clear dataset.", bold_prefix="7. System Logging & Maintenance: ")

    # ==========================================
    # BAB 2: AKTOR SISTEM (SYSTEM ACTORS)
    # ==========================================
    add_heading_styled(doc, "2. AKTOR SISTEM (SYSTEM ACTORS)", level=1)
    add_paragraph_styled(doc, "Dalam aplikasi PADU v1.02, terdapat 2 aktor utama yang berinteraksi dengan sistem:")

    table_actors = doc.add_table(rows=3, cols=3)
    table_actors.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_actors.autofit = False

    headers_act = ["Aktor", "Peran & Deskripsi", "Tanggung Jawab & Tugas Utama"]
    hdr_cells = table_actors.rows[0].cells
    for i, h_text in enumerate(headers_act):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "1F4E78")
        set_cell_margins(hdr_cells[i], top=120, bottom=120)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.name = 'Calibri'
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    actors_data = [
        ("Data Analyst / Auditor Data", "Pengguna tingkat lanjut yang bertanggung jawab mengevaluasi skor kesehatan data dan analisis statistik.", "Menganalisis metrik agregasi, mengevaluasi Quality Score data, menyaring data bermasalah, serta mengunduh laporan audit kesalahan (Error CSV) dan paket data bersih (.ZIP)."),
        ("Operator Data / Administrator Lapangan", "Pengguna operasional yang bertanggung jawab melakukan impor awal data, pencarian data warga, dan maintenance.", "Mengunduh template resmi, melakukan unggah/impor file CSV/XLSX, memantau/membatalkan progres impor, mencari data NIK/KK warga, serta mengelola log & reset dataset.")
    ]

    for row_idx, data in enumerate(actors_data, start=1):
        row_cells = table_actors.rows[row_idx].cells
        bg_color = "F2F2F2" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(data):
            row_cells[col_idx].text = cell_value
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100)
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.name = 'Calibri'
            p.runs[0].font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # BAB 3: MATRIKS DAFTAR USE CASE (UC-01 s/d UC-10)
    # ==========================================
    add_heading_styled(doc, "3. RINGKASAN & MATRIKS USE CASE (UC-01 s/d UC-10)", level=1)
    add_paragraph_styled(doc, "Tabel berikut menyajikan ringkasan 10 Use Case utama yang dikembangkan pada sistem PADU v1.02:")

    table_uc = doc.add_table(rows=11, cols=4)
    table_uc.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_uc.autofit = False

    headers_uc = ["ID Use Case", "Nama Use Case", "Aktor", "Deskripsi Singkat"]
    hdr_cells = table_uc.rows[0].cells
    for i, h_text in enumerate(headers_uc):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "1F4E78")
        set_cell_margins(hdr_cells[i], top=120, bottom=120)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.name = 'Calibri'
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    uc_list = [
        ("UC-01", "Mengimpor Dataset DTSEN", "Operator Data", "Mengunggah file CSV atau XLSX data DTSEN ke dalam aplikasi untuk diproses oleh DuckDB."),
        ("UC-02", "Memantau & Membatalkan Progres Impor", "Operator Data", "Melihat persentase progres ingesti data secara real-time atau menghentikan proses impor via tombol Cancel."),
        ("UC-03", "Mengaudit Kualitas Data (Quality Check)", "Data Analyst", "Melihat ringkasan skor kualitas data, jumlah baris Valid, Warning, dan Critical."),
        ("UC-04", "Melakukan Pencarian & Multi-Filtering Data", "Data Analyst / Operator", "Menyaring data berdasarkan NIK/KK/Nama, Desil, Usia, Pekerjaan, dan Wilayah secara dinamis."),
        ("UC-05", "Meninjau Preview Data Ter-masking (PII Masking)", "Data Analyst / Operator", "Menampilkan rincian baris data dengan penyamaran otomatis pada variabel PII (NIK, Nama, Gaji, Alamat)."),
        ("UC-06", "Mengunduh Template Dataset DTSEN", "Operator Data", "Mengunduh berkas template .csv/.xlsx dengan format header 48 variabel individu & 52 variabel keluarga."),
        ("UC-07", "Mengespor Laporan Data Error (.CSV)", "Data Analyst", "Mengunduh baris data bertanda Critical dalam format CSV untuk perbaikan data di lapangan."),
        ("UC-08", "Mengespor Paket Data Bersih (.ZIP)", "Data Analyst / Operator", "Mengunduh paket terkompresi .zip yang memuat dataset bersih beserta file audit metadata.txt."),
        ("UC-09", "Mengelola Log Pemrosesan Sistem", "Operator Data", "Meninjau live execution logs, mengunduh file log audit, atau membersihkan riwayat log."),
        ("UC-10", "Membersihkan / Reset Dataset", "Operator Data", "Menhapus seluruh dataset yang telah dimuat di DuckDB agar sistem siap digunakan untuk dataset baru.")
    ]

    for row_idx, data in enumerate(uc_list, start=1):
        row_cells = table_uc.rows[row_idx].cells
        bg_color = "F2F2F2" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(data):
            row_cells[col_idx].text = cell_value
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80)
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.name = 'Calibri'
            p.runs[0].font.size = Pt(10)
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================================
    # BAB 4: SPESIFIKASI LENGKAP DETAIL USE CASE (UC-01 s/d UC-10)
    # ==========================================
    add_heading_styled(doc, "4. SPESIFIKASI DETIL USE CASE (FORMAT LENGKAP SRS)", level=1)
    add_paragraph_styled(doc, "Berikut adalah uraian spesifikasi lengkap dan seragam untuk seluruh 10 Use Case (UC-01 s/d UC-10) yang memuat Aktor Utama, Deskripsi, Pre-conditions, Post-conditions, Alur Utama (Main Flow), dan Alur Alternatif (Alternative Flow):")

    all_detailed_ucs = [
        {
            "id": "UC-01",
            "name": "Mengimpor Dataset DTSEN",
            "actor": "Operator Data",
            "desc": "Aktor memasukkan file CSV/XLSX DTSEN ke dalam sistem untuk disimpan dan dianalisis oleh DuckDB.",
            "pre": "Sistem dalam kondisi aktif (http://127.0.0.1:8000/) dan file dataset sesuai format DTSEN 2026.",
            "post": "Data tersimpan di DuckDB engine, skor kualitas data diperbarui, dan KPI statistik ditampilkan pada dashboard.",
            "basic": [
                "1. Aktor membuka halaman beranda PADU.",
                "2. Aktor memilih file CSV/XLSX atau memasukkan local file path.",
                "3. Aktor menekan tombol 'Mulai Impor Data'.",
                "4. Sistem melakukan synonym mapping header kolom dan mengeksekusi script Python DuckDB secara background chunk.",
                "5. Sistem menampilkan progress bar pemrosesan data.",
                "6. Setelah selesai, sistem memperbarui statistik KPI (Total Baris, Total KK, Quality Score, Valid/Warning/Critical)."
            ],
            "alt": [
                "5a. Aktor menekan tombol 'Batal Impor' ➔ Sistem menghentikan pemrosesan subprocess Python dan menghapus data parsial.",
                "4a. Format file tidak valid / terkoorup ➔ Sistem menampilkan pesan error 'Format File Tidak Didukung'."
            ]
        },
        {
            "id": "UC-02",
            "name": "Memantau & Membatalkan Progres Impor",
            "actor": "Operator Data",
            "desc": "Aktor melihat persentase progres ingesti data secara real-time atau menghentikan proses impor via tombol Cancel.",
            "pre": "Proses ingesti data (UC-01) sedang aktif berjalan di latar belakang.",
            "post": "Proses impor selesai 100% atau berhasil dibatalkan tanpa merusak konsistensi data yang ada.",
            "basic": [
                "1. Aktor membuka/melihat panel indikator status pemrosesan impor pada halaman utama.",
                "2. Sistem menampilkan persentase progres dan persentase baris data yang sedang diproses secara real-time.",
                "3. Aktor (opsional) menekan tombol 'Batal Impor' saat pemrosesan berlangsung.",
                "4. Sistem mengirimkan sinyal pembatalan ke subprocess Python DuckDB.",
                "5. Subprocess Python DuckDB menghentikan pembacaan file dan menghapus transaksi memori sementara.",
                "6. Sistem memperbarui antarmuka dengan notifikasi 'Proses Impor Dibatalkan oleh Pengguna'."
            ],
            "alt": [
                "3a. Pemrosesan impor sudah selesai 100% sebelum tombol Batal ditekan ➔ Sistem mengabaikan perintah batal dan menampilkan hasil ingesti selesai."
            ]
        },
        {
            "id": "UC-03",
            "name": "Mengaudit Kualitas Data (Quality Check)",
            "actor": "Data Analyst",
            "desc": "Aktor melihat ringkasan skor kualitas data (Quality Score %), jumlah baris Valid, Warning, dan Critical berdasarkan aturan validasi otomatis.",
            "pre": "Dataset DTSEN sudah berhasil dimuat/diimpor di dalam database DuckDB.",
            "post": "Skor kesehatan data (%) dan rincian jenis kesalahan ditampilkan pada dashboard.",
            "basic": [
                "1. Aktor membuka halaman/dashboard audit kualitas data pada aplikasi PADU.",
                "2. Sistem mengeksekusi query audit kualitas data di DuckDB secara otomatis.",
                "3. Sistem mengklasifikasikan setiap baris data ke dalam kategori Valid, Warning, atau Critical berdasarkan aturan bisnis.",
                "4. Sistem menampilkan skor persen kualitas data (Quality Score %), grafik distribusi status, serta rincian variasi kesalahan data."
            ],
            "alt": [
                "2a. Dataset belum diimpor (database kosong) ➔ Sistem menampilkan pesan 'Belum Ada Data Terimpor. Silakan Impor Dataset Terlebih Dahulu'."
            ]
        },
        {
            "id": "UC-04",
            "name": "Melakukan Pencarian & Multi-Filtering Data",
            "actor": "Data Analyst / Operator Data",
            "desc": "Aktor menyaring dan mencari data warga berdasarkan kata kunci NIK, KK, Nama, maupun kombinasi multi-filter (Desil, Usia, Gaji, Pekerjaan, Wilayah).",
            "pre": "Dataset telah tersedia dan tersimpan dalam sistem DuckDB.",
            "post": "Tabel data dan statistik agregat KPI diperbarui secara real-time sesuai kriteria filter yang dipilih.",
            "basic": [
                "1. Aktor memasukkan kata kunci pencarian (NIK / KK / Nama) pada kolom pencarian.",
                "2. Aktor memilih satu atau beberapa opsi filter kriteria (Desil 1-10, Rentang Usia, Status Bekerja, Wilayah KTP/Keluarga).",
                "3. Aktor menekan tombol 'Cari & Terapkan Filter'.",
                "4. Sistem menyusun query pencarian SQL terparameter di DuckDB dan mengeksekusi query dengan waktu respons < 0.05 detik.",
                "5. Sistem menampilkan daftar data hasil penyaringan beserta metrik statistik KPI khusus data terfilter."
            ],
            "alt": [
                "4a. Data tidak ditemukan sesuai kriteria ➔ Sistem menampilkan tabel kosong dengan pesan 'Data Tidak Ditemukan'."
            ]
        },
        {
            "id": "UC-05",
            "name": "Meninjau Preview Data Ter-masking (PII Masking)",
            "actor": "Data Analyst / Operator Data",
            "desc": "Aktor Menampilkan rincian baris data dengan penyamaran otomatis pada variabel PII (NIK, Nama, Gaji, Alamat).",
            "pre": "Aktor berada di halaman tabel data atau menekan tombol preview detail baris data.",
            "post": "Informasi pribadi sensitif (NIK, Nama, Tanggal Lahir, Nominal Gaji, Alamat) tersaji ter-masking untuk mencegah kebocoran data.",
            "basic": [
                "1. Aktor menekan tombol 'Preview Detail' pada salah satu baris data warga di tabel.",
                "2. Sistem membaca data dari DuckDB dan menerapkan algoritma penyamaran PII (Data Masking Engine).",
                "3. Sistem menyamarkan NIK menjadi 4 digit awal (misal: 3201************), Nama per kata (B*** S******), Tanggal Lahir (1992-**-**), Gaji (Rp 8.xxx.xxx), dan Alamat.",
                "4. Sistem menampilkan pop-up modal detail warga yang aman dari risiko kebocoran privasi."
            ],
            "alt": [
                "2a. Data warga tidak ditemukan/terhapus ➔ Sistem menampilkan error 'Detail Data Tidak Ditemukan'."
            ]
        },
        {
            "id": "UC-06",
            "name": "Mengunduh Template Dataset DTSEN",
            "actor": "Operator Data",
            "desc": "Aktor mengunduh berkas template .csv/.xlsx dengan format header 48 variabel individu & 52 variabel keluarga.",
            "pre": "Sistem dalam kondisi aktif.",
            "post": "Berkas template_dtsen_2026.csv atau .xlsx terunduh ke penyimpanan lokal aktor.",
            "basic": [
                "1. Aktor menekan tombol 'Unduh Template' pada panel impor data.",
                "2. Sistem memanggil controller pembentuk template resmi 48/52 variabel.",
                "3. Browser mengunduh berkas template secara otomatis ke laptop aktor."
            ],
            "alt": [
                "2a. Kegagalan pembuatan file template ➔ Sistem menampilkan pesan 'Gagal Mengunduh Template File'."
            ]
        },
        {
            "id": "UC-07",
            "name": "Mengespor Laporan Data Error (.CSV)",
            "actor": "Data Analyst",
            "desc": "Aktor mengunduh baris data bertanda Critical dalam format CSV untuk perbaikan data di lapangan.",
            "pre": "Proses Quality Check (UC-03) telah dijalankan dan terdapat data berkategori Critical.",
            "post": "Berkas laporan_audit_error.csv terunduh ke komputer lokal.",
            "basic": [
                "1. Aktor menekan tombol 'Ekspor Laporan Error (.CSV)'.",
                "2. Sistem menyaring seluruh baris data berstatus Critical pada database DuckDB.",
                "3. Sistem menambahkan kolom catatan alasan kesalahan (audit_error_log).",
                "4. Sistem mengunduh berkas CSV laporan kesalahan tersebut ke perangkat aktor."
            ],
            "alt": [
                "2a. Tidak ada data berstatus Critical (skor 100% valid) ➔ Sistem menampilkan notifikasi 'Seluruh Data Valid. Tidak Ada Error untuk Diekspor'."
            ]
        },
        {
            "id": "UC-08",
            "name": "Mengespor Paket Data Bersih (.ZIP)",
            "actor": "Data Analyst / Operator Data",
            "desc": "Aktor mengunduh paket terkompresi .zip yang memuat dataset bersih beserta file audit metadata.txt.",
            "pre": "Dataset telah berhasil diimpor dan/atau difilter oleh aktor.",
            "post": "Berkas terkompresi dtsen_export_clean.zip berhasil diunduh.",
            "basic": [
                "1. Aktor menekan tombol 'Ekspor Data Bersih (.ZIP)'.",
                "2. Sistem mengekstrak data hasil olahan/filter di DuckDB menjadi berkas data_dtsen.csv.",
                "3. Sistem membuat berkas metadata.txt yang berisi ringkasan total baris, tanggal ekspor, dan parameter filter yang digunakan.",
                "4. Sistem mengompres kedua berkas tersebut ke dalam bentuk .zip dan memulai proses unduh otomatis."
            ],
            "alt": [
                "2a. Data kosong (0 baris) ➔ Sistem menampilkan peringatan 'Tidak Ada Data untuk Diekspor'."
            ]
        },
        {
            "id": "UC-09",
            "name": "Mengelola Log Pemrosesan Sistem",
            "actor": "Operator Data",
            "desc": "Aktor meninjau live execution logs, mengunduh file log audit, atau membersihkan riwayat log.",
            "pre": "Sistem telah menjalankan aktivitas atau proses ingesti di latar belakang.",
            "post": "Riwayat log dapat ditinjau di layar, diunduh sebagai berkas teks, atau dikosongkan.",
            "basic": [
                "1. Aktor membuka panel 'Log Pemrosesan Sistem'.",
                "2. Sistem menampilkan catatan waktu eksekusi, status query DuckDB, dan pesan aktivitas sistem.",
                "3. Aktor dapat memilih tombol 'Unduh Log' untuk menyimpan file .txt atau tombol 'Bersihkan Log' untuk mengosongkan riwayat."
            ],
            "alt": [
                "3a. Aktor mengonfirmasi 'Bersihkan Log' ➔ Sistem menghapus isi berkas log dan memperbarui tampilan log menjadi kosong."
            ]
        },
        {
            "id": "UC-10",
            "name": "Membersihkan / Reset Dataset",
            "actor": "Operator Data",
            "desc": "Aktor menghapus seluruh dataset yang telah dimuat di DuckDB agar sistem siap digunakan untuk dataset baru.",
            "pre": "Dataset terpasang pada database DuckDB.",
            "post": "Seluruh tabel di DuckDB dikosongkan (0 baris) dan tampilan dashboard kembali ke kondisi awal.",
            "basic": [
                "1. Aktor menekan tombol 'Reset / Hapus Semua Data'.",
                "2. Sistem menampilkan dialog konfirmasi keamanan 'Apakah Anda yakin ingin menghapus seluruh data?'.",
                "3. Aktor menekan konfirmasi 'Ya, Hapus Data'.",
                "4. Sistem mengeksekusi perintah penghapusan tabel pada database DuckDB dan memperbarui antarmuka dashboard ke kondisi kosong."
            ],
            "alt": [
                "3a. Aktor menekan 'Batal' ➔ Sistem membatalkan perintah dan data tetap tersimpan aman."
            ]
        }
    ]

    for uc in all_detailed_ucs:
        add_heading_styled(doc, f"{uc['id']}: {uc['name']}", level=2)
        add_paragraph_styled(doc, uc['actor'], bold_prefix="Aktor Utama: ")
        add_paragraph_styled(doc, uc['desc'], bold_prefix="Deskripsi: ")
        add_paragraph_styled(doc, uc['pre'], bold_prefix="Pre-conditions: ")
        add_paragraph_styled(doc, uc['post'], bold_prefix="Post-conditions: ")
        
        add_paragraph_styled(doc, "Alur Utama (Main Flow):", bold=True)
        for step in uc['basic']:
            add_paragraph_styled(doc, step)
            
        if uc['alt']:
            add_paragraph_styled(doc, "Alur Alternatif (Alternative Flow):", bold=True)
            for alt_step in uc['alt']:
                add_paragraph_styled(doc, alt_step, italic=True)
                
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Save Document
    doc.save(output_path)
    print(f"Dokumen SRS lengkap (UC-01 s/d UC-10) berhasil diperbarui dan disimpan di: {output_path}")

if __name__ == "__main__":
    out_file = r"c:\Users\rasyaad\.gemini\antigravity-ide\scratch\aplikasi-cepat-analytics\Dokumentasi_Use_Case_dan_Keterangan_Aplikasi_PADU_v1.02_Lengkap.docx"
    build_srs_doc(out_file)
